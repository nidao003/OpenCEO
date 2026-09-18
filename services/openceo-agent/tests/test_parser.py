from pathlib import Path

from docx import Document

from app.parser import parse_report


def test_parser_reads_non_uniform_table(tmp_path: Path):
    path = tmp_path / "weekly.docx"
    doc = Document()
    doc.add_paragraph("填报人：peter")
    doc.add_paragraph("汇报周期：2026-09-07 至 2026-09-13")
    table = doc.add_table(rows=3, cols=5)
    headers = ["项目", "工作事项", "本周产出结果", "进度", "状态"]
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    values = ["OpenCEO", "报告解析", "完成第一版", "80%", "正常"]
    for i, value in enumerate(values):
        table.rows[1].cells[i].text = value
    values = ["OpenCEO", "企微接入", "等待企业配置", "50%", "有阻塞"]
    for i, value in enumerate(values):
        table.rows[2].cells[i].text = value
    doc.save(path)

    report = parse_report(path)
    assert report.submitter_name == "peter"
    assert report.period_start.isoformat() == "2026-09-07"
    assert report.period_end.isoformat() == "2026-09-13"
    work = [x for x in report.items if x.item_type == "work"]
    assert len(work) == 2
    assert work[0].raw_project_name == "OpenCEO"
    assert work[0].progress == 80
