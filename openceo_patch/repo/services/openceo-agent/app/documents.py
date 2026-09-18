from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Iterable

from docx import Document

from .config import get_settings


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_docx(path: Path) -> tuple[Path, tempfile.TemporaryDirectory | None]:
    if path.suffix.lower() == ".docx":
        return path, None
    if path.suffix.lower() != ".doc":
        raise ValueError(f"Unsupported Word format: {path.suffix}")

    tmp = tempfile.TemporaryDirectory(prefix="openceo-doc-")
    tmpdir = Path(tmp.name)
    # Always sanitize the conversion filename. Enterprise chat uploads may carry
    # legacy/invalid filename encodings which LibreOffice refuses to open.
    safe_src = tmpdir / "source.doc"
    shutil.copy2(path, safe_src)
    cmd = [get_settings().libreoffice_bin, "--headless", "--convert-to", "docx", "--outdir", str(tmpdir), str(safe_src)]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
    converted = tmpdir / "source.docx"
    if proc.returncode != 0 or not converted.exists():
        tmp.cleanup()
        raise RuntimeError(f"LibreOffice conversion failed: {proc.stderr or proc.stdout}")
    return converted, tmp


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def iter_docx_content(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    doc = Document(str(path))
    paragraphs = [_clean(p.text) for p in doc.paragraphs if _clean(p.text)]
    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)
    return paragraphs, tables


def flatten_text(paragraphs: Iterable[str], tables: list[list[list[str]]]) -> str:
    blocks = [*paragraphs]
    for table in tables:
        for row in table:
            blocks.append(" | ".join(row))
    return "\n".join(blocks)
