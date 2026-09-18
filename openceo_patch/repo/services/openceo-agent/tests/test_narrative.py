from pathlib import Path
from docx import Document
from app.parser import parse_report


def test_narrative_metrics_are_preserved(tmp_path: Path):
    path = tmp_path / "narrative.docx"
    doc = Document()
    doc.add_paragraph("填报人：张三")
    doc.add_paragraph("本周重点工作")
    doc.add_paragraph("新增1个产品。")
    doc.add_paragraph("开通5个站安检机上广告牌点位。")
    doc.add_paragraph("项目经营情况：9月执行额120万元，签约额200万元。")
    doc.save(path)
    report = parse_report(path)
    texts = [i.content for i in report.items]
    assert any("新增1个产品" in t for t in texts)
    assert any("执行额120万元" in t for t in texts)
